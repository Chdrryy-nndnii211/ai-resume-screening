"""screening.py - core engine for the AI Resume Screening & Shortlisting system.

Pipeline: read file -> (optional) mask PII -> extract skills / years / education
-> semantic similarity (Sentence-BERT) -> hybrid weighted score -> ranked table.
"""
import io
import re
from datetime import datetime

import numpy as np
import pandas as pd

MODEL_NAME = "sentence-transformers/all-MiniLM-L6-v2"
_model = None


def get_model():
    """Load the Sentence-BERT model once (lazy import keeps helpers lightweight)."""
    global _model
    if _model is None:
        from sentence_transformers import SentenceTransformer
        _model = SentenceTransformer(MODEL_NAME)
    return _model


# ----------------------------------------------------------------------------
# 1. Skill taxonomy: canonical name -> aliases. Extend this freely for your domain.
# ----------------------------------------------------------------------------
SKILLS = {
    # programming
    "python": [], "java": [], "javascript": ["js"], "typescript": [],
    "c++": ["cpp"], "c#": [], "php": [], "ruby": [], "kotlin": [], "swift": [],
    "r programming": ["r language"], "matlab": [], "scala": [], "html": ["html5"],
    "css": ["css3"], "bash": ["shell scripting"],
    # web / frameworks
    "react": ["reactjs", "react.js"], "angular": ["angularjs"], "vue": ["vue.js", "vuejs"],
    "node.js": ["nodejs", "node js"], "django": [], "flask": [], "fastapi": [],
    "spring boot": ["springboot"], ".net": ["dotnet", "asp.net"],
    # data & databases
    "sql": ["t-sql", "pl/sql"], "mysql": [], "postgresql": ["postgres"],
    "mongodb": ["mongo"], "oracle": [], "nosql": [], "excel": ["ms excel", "microsoft excel"],
    "power bi": ["powerbi"], "tableau": [], "data analysis": ["data analytics"],
    "data visualization": ["data visualisation"], "etl": [], "hadoop": [],
    "spark": ["pyspark", "apache spark"], "kafka": [],
    # machine learning
    "machine learning": ["ml"], "deep learning": [],
    "nlp": ["natural language processing"], "computer vision": [],
    "tensorflow": [], "pytorch": [], "keras": [], "scikit-learn": ["sklearn", "scikit learn"],
    "pandas": [], "numpy": [], "statistics": ["statistical analysis"],
    "llm": ["llms", "large language models"],
    # cloud / devops
    "aws": ["amazon web services"], "azure": [], "gcp": ["google cloud"],
    "docker": [], "kubernetes": ["k8s"], "git": ["github", "gitlab"],
    "ci/cd": ["jenkins"], "linux": [], "rest api": ["restful", "rest apis", "restful api"],
    "microservices": [], "agile": ["scrum"], "jira": [],
    "networking": ["network administration"], "technical support": ["help desk", "helpdesk"],
    # business / management
    "project management": [], "leadership": [], "communication": [],
    "customer service": [], "sales": [], "marketing": [], "digital marketing": [],
    "seo": [], "negotiation": [], "business development": [],
    "budgeting": ["budget management"], "financial reporting": [],
    "accounting": [], "accounts payable": [], "accounts receivable": [],
    "general ledger": [], "auditing": ["audit"], "taxation": ["tax preparation"],
    "payroll": [], "quickbooks": [], "sap": [], "tally": [],
    # HR
    "recruitment": ["recruiting", "talent acquisition"], "onboarding": [],
    "employee relations": [], "performance management": [], "hris": [],
    "benefits administration": [], "training": ["training and development"],
    # engineering / operations
    "autocad": [], "solidworks": [], "cad": [], "plc": [], "six sigma": [],
    "quality control": ["quality assurance"], "project planning": [],
    "manufacturing": [], "supply chain": [], "inventory management": [],
    # kitchen / hospitality / healthcare
    "menu planning": [], "food safety": ["haccp"], "catering": [],
    "kitchen management": [], "patient care": [],
}
ALL_SKILLS = sorted(SKILLS)


def _build_patterns():
    pats = {}
    for canon, aliases in SKILLS.items():
        names = sorted({canon, *aliases}, key=len, reverse=True)
        body = "|".join(re.escape(n) for n in names)
        pats[canon] = re.compile(r"(?<![a-z0-9+#.])(?:" + body + r")(?![a-z0-9+#])")
    return pats


_PATTERNS = _build_patterns()


def extract_skills(text):
    """Return the set of canonical skills mentioned in `text`."""
    t = str(text).lower()
    return {canon for canon, pat in _PATTERNS.items() if pat.search(t)}


# ----------------------------------------------------------------------------
# 2. Experience, education, extras
# ----------------------------------------------------------------------------
def extract_years(text):
    """Estimate years of experience: explicit '5 years' claims, else date ranges."""
    t = str(text).lower()
    explicit = [int(x) for x in re.findall(r"(\d{1,2})\+?\s*(?:years|yrs)", t)]
    explicit = [x for x in explicit if 0 < x <= 40]
    if explicit:
        return float(max(explicit))
    now, total = datetime.now().year, 0.0
    pattern = (r"((?:19|20)\d{2})\s*(?:-|to|\u2013|\u2014)\s*(?:\d{1,2}/)?"
               r"((?:19|20)\d{2}|present|current|now)")
    for start, end in re.findall(pattern, t):
        s = int(start)
        e = now if end in ("present", "current", "now") else int(end)
        if 0 <= e - s <= 40:
            total += e - s
    return float(min(total, 40))


EDU_REQ = {"Any": 0, "Diploma": 1, "Bachelor": 2, "Master": 3, "PhD": 4}


def education_level(text):
    """0 = none found, 1 = diploma, 2 = bachelor, 3 = master, 4 = PhD."""
    t = str(text).lower()
    if re.search(r"\b(ph\.?d|doctorate|doctoral)\b", t):
        return 4
    if re.search(r"\b(master'?s?|mba|m\.?tech|msc|m\.sc)\b", t):
        return 3
    if re.search(r"\b(bachelor'?s?|b\.?tech|bsc|b\.sc|bba|bca|b\.e\.)", t):
        return 2
    if re.search(r"\b(diploma|associate'?s?)\b", t):
        return 1
    return 0


_EXTRA_RE = re.compile(r"certif|award|publication|patent|project|hackathon|open[- ]source")

# ----------------------------------------------------------------------------
# 3. Bias mitigation: blind screening
# ----------------------------------------------------------------------------
def mask_pii(text):
    """Mask contact details and gendered words. (Names need NER - see README.)"""
    t = str(text)
    t = re.sub(r"[\w.+-]+@[\w-]+\.[\w.-]+", "[EMAIL]", t)
    t = re.sub(r"(?:https?://|www\.)\S+|linkedin\.com/\S+|github\.com/\S+", "[URL]", t, flags=re.I)
    t = re.sub(r"(?<!\d)(?:\+?\d{1,3}[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}(?!\d)", "[PHONE]", t)
    t = re.sub(r"\b(he|she|his|hers?|him|mr|mrs|ms|miss)\b\.?", "[X]", t, flags=re.I)
    return t


# ----------------------------------------------------------------------------
# 4. Semantic similarity with Sentence-BERT (chunked so long resumes are covered)
# ----------------------------------------------------------------------------
def chunk_text(text, size=120, overlap=20):
    words = str(text).split()
    if not words:
        return [""]
    return [" ".join(words[i:i + size])
            for i in range(0, max(len(words) - overlap, 1), size - overlap)]


def embed_resume_chunks(texts, batch_size=64, show_progress=False):
    """Return one (n_chunks, dim) array per resume. Reuse it across many JDs."""
    model = get_model()
    all_chunks, counts = [], []
    for t in texts:
        c = chunk_text(t)
        all_chunks.extend(c)
        counts.append(len(c))
    embs = model.encode(all_chunks, batch_size=batch_size, normalize_embeddings=True,
                        convert_to_numpy=True, show_progress_bar=show_progress)
    out, pos = [], 0
    for n in counts:
        out.append(embs[pos:pos + n])
        pos += n
    return out


def jd_embedding(jd):
    e = get_model().encode(chunk_text(jd), normalize_embeddings=True,
                           convert_to_numpy=True, show_progress_bar=False)
    v = e.mean(axis=0)
    return v / (np.linalg.norm(v) + 1e-9)


def semantic_raw(jd_vec, chunk_embs, top_k=3):
    """Per resume: mean of the top-k chunk similarities to the JD."""
    return np.array([float(np.sort(e @ jd_vec)[-top_k:].mean()) for e in chunk_embs])


SIM_LOW, SIM_HIGH = 0.10, 0.60   # raw cosine range mapped to 0..1


# ----------------------------------------------------------------------------
# 5. Hybrid scoring + ranking
# ----------------------------------------------------------------------------
DEFAULT_WEIGHTS = {"semantic": 0.30, "skills": 0.35, "experience": 0.20,
                   "education": 0.10, "extras": 0.05}


def label_for(score):
    return "Strong Match" if score >= 75 else "Moderate Match" if score >= 50 else "Weak Match"


def score_candidates(resumes, jd, names=None, must=None, nice=None, min_years=0,
                     min_education="Any", weights=None, blind=False, chunk_embs=None):
    """Score every resume against the JD and return a ranked DataFrame."""
    n = len(resumes)
    if blind:
        names = [f"Candidate {i + 1}" for i in range(n)]
    elif names is None:
        names = [f"Candidate {i + 1}" for i in range(n)]

    w = {**DEFAULT_WEIGHTS, **(weights or {})}
    total = sum(w.values()) or 1.0
    w = {k: v / total for k, v in w.items()}

    texts = [mask_pii(r) for r in resumes] if blind else list(resumes)
    if blind or chunk_embs is None:          # embeddings must match the (masked) text
        chunk_embs = embed_resume_chunks(texts)
    raw = semantic_raw(jd_embedding(jd), chunk_embs)
    semantic = np.clip((raw - SIM_LOW) / (SIM_HIGH - SIM_LOW), 0, 1)

    jd_skills = extract_skills(jd)
    must = {s.lower() for s in must} if must else set(jd_skills)
    nice = {s.lower() for s in nice} if nice else (jd_skills - must)
    denom = 2 * len(must) + len(nice)
    req_edu = EDU_REQ.get(min_education, 0)

    rows = []
    for i, t in enumerate(texts):
        sk = extract_skills(t)
        skill = (2 * len(must & sk) + len(nice & sk)) / denom if denom else float(semantic[i])
        yrs = extract_years(t)
        exp = 1.0 if min_years <= 0 else min(1.0, yrs / min_years)
        edu = 1.0 if req_edu == 0 else min(1.0, education_level(t) / req_edu)
        extras = min(1.0, len(_EXTRA_RE.findall(t.lower())) / 5)
        final = 100 * (w["semantic"] * semantic[i] + w["skills"] * skill +
                       w["experience"] * exp + w["education"] * edu + w["extras"] * extras)
        rows.append({
            "idx": i, "candidate": names[i], "final_score": round(final, 1),
            "match": label_for(final), "semantic": round(float(semantic[i]), 3),
            "skills": round(skill, 3), "experience": round(exp, 3),
            "education": round(edu, 3), "extras": round(extras, 3),
            "years_exp": yrs,
            "matched_skills": ", ".join(sorted((must | nice) & sk)),
            "missing_skills": ", ".join(sorted((must | nice) - sk)),
        })
    return (pd.DataFrame(rows).sort_values("final_score", ascending=False)
            .reset_index(drop=True))


# ----------------------------------------------------------------------------
# 6. Reading uploaded files (PDF / DOCX / TXT)
# ----------------------------------------------------------------------------
def read_file(name, data):
    """Extract text from an uploaded file's bytes."""
    ext = name.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n".join((p.extract_text() or "") for p in pdf.pages)
    if ext == "docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    return data.decode("utf-8", errors="ignore")
